#!/usr/bin/env python3
"""Google Sheets sync over the operator's existing rclone Drive OAuth token.

One consistent approach for every Google integration in the bounty-VRF
family: the rclone configuration file in the owning profile already carries
a Google OAuth token (refresh token + access token, drive scope). This
script refreshes it against the well-known rclone Google OAuth client
(public client id; the public client secret is recovered from the
constant embedded in the rclone binary, de-obfuscated with rclone's
public AES-CTR scheme) and then talks to the Drive REST API directly:

  read   export the spreadsheet to .xlsx (Drive v3 /files/{id}/export)
  write  in-place content update (upload/drive/v3/files/{id}, multipart,
         xlsx body, original mimeType preserved) + read-back
  log    targeted cell updates in the "Challenge Log" tab keyed by Case id

No browser UI. No Sheets API (it is disabled in the rclone OAuth project);
the Drive export/upload round-trip covers both directions. The token cache
is written next to the rclone config with mode 0600 and never printed.

Environment overrides (all optional):
  BBOPS_RC_CONF      path to rclone.conf (default: --profile-dir/data/rclone.conf)
  BBOPS_TOKEN_CACHE  path of the token cache file
  BBOPS_SHEET_ID     default spreadsheet id
  BBOPS_DRIVE_API    Drive API base (https://www.googleapis.com by default;
                     override for tests)
  BBOPS_OAUTH_URL    OAuth token endpoint (https://oauth2.googleapis.com/token)
  BBOPS_UPLOAD_API   upload base (https://www.googleapis.com by default)
"""
import argparse
import base64
import hashlib
import io
import json
import os
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
from datetime import datetime, timezone
from pathlib import Path

CLIENT_ID = "202264815644.apps.googleusercontent.com"
# rclone's public Google Drive client secret, encrypted with rclone's
# public obscure scheme (AES-CTR, key below). Same constant as the
# rclone v1.73.5 drive backend embeds.
ENCRYPTED_CLIENT_SECRET = (
    "eX8GpZTVx3vxMWVkuuBdDWmAUE6rGhTwVrvG9GhllYccSdj2-mvHVg")
_OBSCURE_KEY = bytes(
    [0x9C, 0x93, 0x5B, 0x48, 0x73, 0x0A, 0x55, 0x4D,
     0x6B, 0xFD, 0x7C, 0x63, 0xC8, 0x86, 0xA9, 0x2B,
     0xD3, 0x90, 0x19, 0x8E, 0xB8, 0x12, 0x8A, 0xFB,
     0xF4, 0xDE, 0x16, 0x2B, 0x8B, 0x95, 0xF6, 0x38])
XLSX_MIME = ("application/vnd.openxmlformats-officedocument."
             "spreadsheetml.sheet")
SHEET_MIME = "application/vnd.google-apps.spreadsheet"
LOG_TAB = "Challenge Log"
LOG_HEADERS = ["Case", "Program", "Status", "Repository / commit",
               "Mithril evaluation", "Local validation",
               "Confirmed vulnerabilities", "Submitted reports",
               "Desktop result", "Model qualification", "Scope gate",
               "Evidence", "Next step", "Date JST"]

DEFAULT_SHEET_ID = (
    "1gJ-CfDviZZltpc2qIlsW5hxs386E9fCWhkS0vR_50yo")


def now_iso():
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def sha256_bytes(b):
    return hashlib.sha256(b).hexdigest()


def reveal_rclone_secret(x):
    """De-obfuscate a rclone obscure: value (AES-CTR, public key)."""
    try:
        from cryptography.hazmat.primitives.ciphers import (Cipher,
                                                            algorithms,
                                                            modes)
    except ImportError:
        return None
    raw = base64.urlsafe_b64decode((x + "=" * (-len(x) % 4)).encode())
    if len(raw) < 17:
        return None
    dec = Cipher(algorithms.AES(_OBSCURE_KEY),
                 modes.CTR(raw[:16])).decryptor()
    return (dec.update(raw[16:]) + dec.finalize()).decode("utf-8",
                                                          "replace")


CLIENT_SECRET = reveal_rclone_secret(ENCRYPTED_CLIENT_SECRET)


class DriveError(Exception):
    def __init__(self, code, body):
        super().__init__("drive %s: %s" % (code, body[:300]))
        self.code = code
        self.body = body


class Client:
    """Minimal Drive API client with token refresh + quota backoff."""

    def __init__(self, rclone_conf, token_cache):
        self.rclone_conf = Path(rclone_conf)
        self.token_cache = Path(token_cache)
        self.api = os.environ.get("BBOPS_DRIVE_API",
                                  "https://www.googleapis.com")
        self.upload = os.environ.get("BBOPS_UPLOAD_API",
                                     "https://www.googleapis.com")
        self.oauth = os.environ.get("BBOPS_OAUTH_URL",
                                    "https://oauth2.googleapis.com/token")
        self._token = None
        self._expiry = 0.0

    # -- token handling -------------------------------------------------
    def _rclone_token(self):
        text = self.rclone_conf.read_text()
        m = re.search(r"^\s*token\s*=\s*(\{.*\})\s*$", text, re.M)
        if not m:
            raise DriveError("conf", "no drive token in %s"
                             % self.rclone_conf)
        return json.loads(m.group(1))

    def _refresh(self, refresh_token):
        body = urllib.parse.urlencode({
            "grant_type": "refresh_token",
            "refresh_token": refresh_token,
            "client_id": CLIENT_ID,
            "client_secret": CLIENT_SECRET,
        }).encode()
        req = urllib.request.Request(
            self.oauth, data=body,
            headers={"Content-Type": "application/x-www-form-urlencoded"})
        try:
            with urllib.request.urlopen(req, timeout=30) as r:
                tok = json.loads(r.read())
        except urllib.error.HTTPError as e:
            raise DriveError(e.code, e.read().decode("utf-8", "replace"))
        if "access_token" not in tok:
            raise DriveError("token", str(tok))
        tok["_expires_at"] = time.time() + int(
            tok.get("expires_in", 3500)) - 60
        return tok

    def access_token(self):
        if self._token and time.time() < self._expiry - 60:
            return self._token["access_token"]
        tok = None
        try:
            cached = json.loads(self.token_cache.read_text())
            if time.time() < cached.get("_expires_at", 0) - 60:
                tok = cached
        except (OSError, ValueError):
            pass
        if tok is None:
            tok = self._refresh(self._rclone_token()["refresh_token"])
        self.token_cache.parent.mkdir(parents=True, exist_ok=True)
        tmp = self.token_cache.with_suffix(".tmp")
        tmp.write_text(json.dumps(tok))
        os.chmod(tmp, 0o600)
        os.replace(tmp, self.token_cache)
        self._token = tok
        self._expiry = tok.get("_expires_at", 0.0)
        return tok["access_token"]

    # -- http -----------------------------------------------------------
    def call(self, url, method="GET", data=None, headers=None, tries=8):
        last = (0, b"")
        for attempt in range(tries):
            req = urllib.request.Request(
                url, data=data, method=method,
                headers=dict(headers or {}, Authorization="Bearer "
                                   + self.access_token()))
            try:
                with urllib.request.urlopen(req, timeout=180) as r:
                    return r.status, r.read()
            except urllib.error.HTTPError as e:
                last = (e.code, e.read())
                if e.code == 403 and b"Quota" in last[1] \
                        and attempt < tries - 1:
                    time.sleep(20 + attempt * 10)
                    continue
                raise DriveError(*last) from None
        raise DriveError(*last)

    # -- sheet operations -------------------------------------------------
    def export(self, sheet_id):
        _, data = self.call("%s/drive/v3/files/%s/export?mimeType=%s"
                            % (self.api, sheet_id,
                               urllib.parse.quote(XLSX_MIME)))
        return data

    def update(self, sheet_id, xlsx_bytes, name=None):
        boundary = "bops" + uuid.uuid4().hex
        meta = json.dumps({"mimeType": SHEET_MIME}
                          | ({"name": name} if name else {})).encode()
        body = b"".join([
            ("--%s\r\nContent-Type: application/json; charset=UTF-8"
             "\r\n\r\n" % boundary).encode(), meta,
            ("\r\n--%s\r\nContent-Type: %s\r\n\r\n"
             % (boundary, XLSX_MIME)).encode(), xlsx_bytes,
            ("\r\n--%s--\r\n" % boundary).encode(),
        ])
        code, _ = self.call(
            "%s/upload/drive/v3/files/%s?uploadType=multipart"
            % (self.upload, sheet_id),
            method="PATCH", data=body,
            headers={"Content-Type": "multipart/related; boundary=%s"
                     % boundary})
        return code == 200

    def log_row(self, sheet_id, case_id, sets):
        """Update cells in the LOG_TAB row whose column A equals case_id.

        sets: {header: value}. Returns (written_cells, readback_ok).
        """
        import openpyxl  # optional dependency, needed only for log/write
        wb = openpyxl.load_workbook(io.BytesIO(self.export(sheet_id)))
        if LOG_TAB not in wb.sheetnames:
            raise DriveError("tabs", "no %r tab; have %s"
                             % (LOG_TAB, wb.sheetnames))
        ws = wb[LOG_TAB]
        header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        col = {str(h).strip(): i + 1 for i, h in enumerate(header)
               if h is not None}
        missing = [k for k in sets if k not in col]
        if missing:
            raise DriveError("headers", "unknown columns %s; have %s"
                             % (missing, sorted(col)))
        target = None
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(r, 1).value or "").strip() == case_id:
                target = r
                break
        if target is None:
            raise DriveError("row", "no row with Case %r" % case_id)
        for k, v in sets.items():
            ws.cell(target, col[k]).value = v
        buf = io.BytesIO()
        wb.save(buf)
        self.update(sheet_id, buf.getvalue())
        # read-back verification
        wb2 = openpyxl.load_workbook(
            io.BytesIO(self.export(sheet_id)))
        ws2 = wb2[LOG_TAB]
        ok = all(
            str(ws2.cell(target, col[k]).value or "") == str(v)
            for k, v in sets.items())
        return list(sets), ok


def parse_args(argv):
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--profile-dir",
                   default=os.environ.get(
                       "BBOPS_PROFILE_DIR",
                       os.path.expanduser(
                           "~/.hermes/profiles/mithril-bounty-vrf-ops")))
    p.add_argument("--rclone-conf")
    p.add_argument("--token-cache")
    p.add_argument("--sheet-id",
                   default=os.environ.get("BBOPS_SHEET_ID",
                                          DEFAULT_SHEET_ID))
    sub = p.add_subparsers(dest="cmd", required=True)
    sub.add_parser("refresh")
    r = sub.add_parser("read")
    r.add_argument("--out", required=True)
    w = sub.add_parser("write")
    w.add_argument("--xlsx", required=True)
    lg = sub.add_parser("log")
    lg.add_argument("--case", required=True)
    lg.add_argument("--set", action="append", default=[],
                    metavar="COL=VALUE",
                    help="header=value (repeatable)")
    return p.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)
    profile = Path(args.profile_dir).expanduser()
    conf = Path(args.rclone_conf) if args.rclone_conf else \
        Path(os.environ.get(
            "BBOPS_RC_CONF", profile / "data" / "rclone.conf"))
    cache = Path(args.token_cache) if args.token_cache else \
        Path(os.environ.get(
            "BBOPS_TOKEN_CACHE", profile / "data" / "sheets_token.json"))
    client = Client(conf, cache)
    try:
        if args.cmd == "refresh":
            at = client.access_token()
            scope = (client._token or {}).get("scope", "unknown")
            print("TOKEN\tok len=%d scope=%s" % (len(at), scope))
            return 0
        if args.cmd == "read":
            data = client.export(args.sheet_id)
            out = Path(args.out)
            tmp = out.with_suffix(out.suffix + ".tmp")
            tmp.write_bytes(data)
            os.replace(tmp, out)
            print("READ\t%s bytes=%d sha256=%s"
                  % (out, len(data), sha256_bytes(data)))
            return 0
        if args.cmd == "write":
            data = Path(args.xlsx).read_bytes()
            client.update(args.sheet_id, data)
            back = client.export(args.sheet_id)
            print("WRITE\tok uploaded=%d readback_bytes=%d readback_sha=%s"
                  % (len(data), len(back), sha256_bytes(back)))
            return 0
        if args.cmd == "log":
            sets = {}
            for item in args.set:
                if "=" not in item:
                    print("LOG\tFAILED bad --set %r" % item)
                    return 2
                k, v = item.split("=", 1)
                sets[k.strip()] = v
            written, ok = client.log_row(args.sheet_id, args.case, sets)
            print("LOG\t%s case=%s cells=%s readback=%s"
                  % ("OK" if ok else "READBACK-MISMATCH", args.case,
                     len(written), "match" if ok else "mismatch"))
            return 0 if ok else 1
    except DriveError as e:
        print("SHEETS\tFAILED %s" % e)
        return 1
    except OSError as e:
        print("SHEETS\tFAILED %s" % e)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
